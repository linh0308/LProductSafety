import hashlib
import io
import json
import mimetypes
import secrets
from datetime import date
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Query
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, update, delete, func
from sqlalchemy.orm import Session
from backend.store import Account, LoginSession, Record, Evidence, Audit, ImportLog, Setting, get_db, log, now
from backend.catalog import MODULES, catalog
from backend.excel_import import FILES, ROOT, import_sources
from backend.operations import derived_status, xrf_evaluation, workbook_conflicts
from login import current_user, admin_user, profile, ROLES

router=APIRouter(prefix='/api')
UPLOADS=ROOT/'data'/'uploads'
ACTIONS=['View','Create','Edit','Delete','Upload','Approve','Close']


def permissions(user, db):
    if user.role=='Admin': return {m:{a:True for a in ACTIONS} for m in MODULES}
    custom=db.get(Setting,'permissions')
    if custom and user.role in custom.value:
        return custom.value[user.role]
    result={}
    for key,config in MODULES.items():
        write=user.role in ('QA Manager','Product Safety Engineer') or user.role=='IQC' and key in ('xrf-iqc','xrf-plan') or user.role=='OQC' and key in ('xrf-oqc','change-control') or user.role=='Supplier Quality' and key in ('suppliers','fmd','reports','material-declarations','documents')
        write=write and config['group']!='settings'
        result[key]={a:(True if a=='View' else user.role=='QA Manager' if a in ('Approve','Close','Delete') else write) for a in ACTIONS}
    return result


def permit(user,db,module,action='View'):
    if module not in MODULES: raise HTTPException(404,'Không có module này.')
    if not permissions(user,db).get(module,{}).get(action,False): raise HTTPException(403,'Bạn không có quyền '+action+' cho module này.')


def record_or_404(db,record_id):
    row=db.get(Record,record_id)
    if not row or row.archived: raise HTTPException(404,'Không tìm thấy hồ sơ.')
    return row


def all_visible(db,user):
    allowed=[m for m,p in permissions(user,db).items() if p.get('View')]
    return db.scalars(select(Record).where(Record.archived==False,Record.module.in_(allowed))).all()


def standards(db):
    return db.scalars(select(Record).where(Record.module=='xrf-standard',Record.archived==False)).all()


def serialize(row,limits):
    data=dict(row.data)
    if row.module=='materials':
        data['usage_status']=data.get('usage_status') or 'Đang sử dụng'
        data['status']=data.get('status') or 'Active'
    mat_display_status = data.get('status') or data.get('usage_status') or 'Đang sử dụng'
    return {'id':row.id,'module':row.module,'data':data,'version':row.version,'updated_at':row.updated_at.isoformat(),
            'display_status':mat_display_status if row.module=='materials' else derived_status(row.module,row.data,limits)}


@router.get('/catalog')
def get_catalog(user: Account=Depends(current_user),db: Session=Depends(get_db)):
    return {**catalog(),'permissions':permissions(user,db),'roles':ROLES}


def filtered(rows,module,q='',status='',start='',end='',material='',dossier='',project='',category=''):
    output=[]
    for row in rows:
        if module and row['module']!=module: continue
        d=row['data']
        if q and q.casefold() not in json.dumps(d,ensure_ascii=False).casefold(): continue
        if project and project.casefold().replace('-', ' ').strip() not in str(d.get('project','')).casefold().replace('-', ' '): continue
        if category and category.casefold() not in str(d.get('category','')).casefold(): continue
        if status and row['display_status']!=status: continue
        if dossier and d.get('dossier_status')!=dossier: continue
        if material and d.get('material_code')!=material: continue
        dt=str(d.get('test_date') or d.get('expiry_date') or d.get('due_date') or d.get('issue_date') or '')
        if start and (not dt or dt<start): continue
        if end and (not dt or dt>end): continue
        output.append(row)
    return output


@router.get('/records')
def records(module: str, q: str='',status: str='',start: str='',end: str='', material: str='',dossier: str='',project: str='',category: str='',sort: str='updated_at',direction: str='desc',page: int=Query(1,ge=1),size: int=Query(20,ge=1,le=2000),user: Account=Depends(current_user),db: Session=Depends(get_db)):
    permit(user,db,module)
    limits=standards(db)
    rows=[serialize(r,limits) for r in db.scalars(select(Record).where(Record.module==module,Record.archived==False)).all()]
    rows=filtered(rows,module,q,status,start,end,material,dossier,project,category)
    rows.sort(key=lambda r:str(r['updated_at'] if sort=='updated_at' else r['data'].get(sort,'')).casefold(),reverse=direction=='desc')
    return {'items':rows[(page-1)*size:page*size],'total':len(rows),'page':page,'size':size}


@router.get('/search')
def search(q: str=Query('',max_length=200),user: Account=Depends(current_user),db: Session=Depends(get_db)):
    if len(q.strip())<2: return []
    limits=standards(db)
    return [serialize(r,limits) for r in all_visible(db,user) if q.casefold() in json.dumps(r.data,ensure_ascii=False).casefold()][:50]


INSPECTION_SETS={'inspection-results':('reports','oqc-reports','xrf-iqc','xrf-oqc','change-control'), 'inspection-plans':('test-plan','xrf-plan')}


@router.get('/inspections/{collection}')
def inspections(collection:str,q:str='',stage:str='',method:str='',result:str='',validity:str='',expiry_days:int|None=Query(None,ge=0,le=3650),page:int=Query(1,ge=1),size:int=Query(20,ge=1,le=2000),user:Account=Depends(current_user),db:Session=Depends(get_db)):
    modules=INSPECTION_SETS.get(collection)
    if not modules: raise HTTPException(404)
    allowed=[m for m in modules if permissions(user,db).get(m,{}).get('View')]
    if not allowed: raise HTTPException(403,'Bạn không có quyền xem dữ liệu kiểm nghiệm.')
    limits=standards(db); rows=[]
    for record in db.scalars(select(Record).where(Record.archived==False,Record.module.in_(allowed))).all():
        row=serialize(record,limits); d=row['data']; module=row['module']
        row['stage']=d.get('inspection_stage') or {'reports':'IQC','oqc-reports':'OQC','xrf-iqc':'IQC','xrf-oqc':'OQC','change-control':'Khác'}.get(module) or (d.get('test') if d.get('test') in ('IQC','OQC') else 'Chưa xác định')
        row['method']='XRF' if module.startswith('xrf-') or module=='change-control' else 'Lab / Bên thứ ba'
        raw=str(d.get('result','')).strip().upper()
        row['test_result']=('PASS' if row['display_status']=='Pass' else 'FAIL' if row['display_status']=='NG' else 'Chưa đánh giá') if row['method']=='XRF' else ('PASS' if raw=='PASS' else 'FAIL' if raw in ('FAIL','NG') else 'Chưa đánh giá')
        try:
            days=(date.fromisoformat(str(d.get('expiry_date')))-date.today()).days
            row['validity']='Hết hạn' if days<0 else 'Sắp hết hạn' if days<=90 else 'Còn hạn'
        except (ValueError,TypeError): row['validity']='Chưa có hạn'
        if expiry_days is not None:
            try: remaining=(date.fromisoformat(str(d.get('expiry_date')))-date.today()).days
            except (ValueError,TypeError): continue
            if not 0<=remaining<=expiry_days: continue
        if q and q.casefold() not in json.dumps(d,ensure_ascii=False).casefold(): continue
        if stage and stage!=row['stage'] or method and method!=row['method'] or result and result!=row['test_result'] or validity and validity!=row['validity']: continue
        rows.append(row)
    rows.sort(key=lambda r:(r['updated_at'],r['id']),reverse=True)
    total=len(rows);page=min(page,max(1,(total+size-1)//size))
    return {'items':rows[(page-1)*size:page*size],'total':total,'page':page,'size':size}


@router.get('/records/{record_id}')
def detail(record_id:int,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    row=record_or_404(db,record_id); permit(user,db,row.module)
    limits=standards(db); result=serialize(row,limits)
    result['evidence']=[{'id':e.id,'name':e.name,'mime':e.mime,'size':e.size,'checksum':e.checksum,'uploader':e.uploader,'created_at':e.created_at.isoformat()} for e in db.scalars(select(Evidence).where(Evidence.record_id==row.id)).all()]
    result['history']=[{'id':a.id,'actor':a.actor,'action':a.action,'at':a.created_at.isoformat(),'before':a.before,'after':a.after} for a in db.scalars(select(Audit).where(Audit.record_id==row.id).order_by(Audit.id.desc())).all()]
    if row.module in ('xrf-iqc','xrf-oqc','change-control'): result['evaluation']=xrf_evaluation(row.data,limits)
    if row.module=='materials':
        code=row.data['material_code']; aliases={code}
        visible=all_visible(db,user)
        for r in visible:
            if r.module=='xrf-plan' and r.data.get('material_code')==code and r.data.get('old_code'):
                aliases.add(r.data['old_code'])
        supplier_names={r.data.get('supplier') for r in visible if r.data.get('material_code') in aliases and r.module in ('materials','bom','reports','fmd')}
        supplier_names.update(r.data.get('supplier') for r in visible if r.module=='suppliers' and aliases.intersection(str(r.data.get('material_codes','')).splitlines()))
        supplier_names.add(row.data.get('supplier'));supplier_names.discard(None);supplier_names.discard('')
        related=[r for r in visible if r.id!=row.id and (r.data.get('material_code') in aliases or r.module=='suppliers' and r.data.get('supplier') in supplier_names)]
        result['supplier_names']=sorted(supplier_names)
        result['related']=[serialize(r,limits) for r in related]
        related_ids=[r.id for r in related]
        attachments={}
        if related_ids:
            for evidence in db.scalars(select(Evidence).where(Evidence.record_id.in_(related_ids))).all():
                attachments.setdefault(evidence.record_id,[]).append({'id':evidence.id,'name':evidence.name})
        for item in result['related']:
            item['files']=attachments.get(item['id'],[])
        result['aliases']=sorted(aliases)
        result['conflicts']=[c for c in workbook_conflicts(visible) if c['material_code']==code]
    if row.module=='suppliers':
        name=row.data.get('supplier')
        visible=all_visible(db,user)
        supplied=[r for r in visible if name and r.id!=row.id and r.data.get('supplier')==name]
        codes={r.data.get('material_code') for r in supplied if r.module in ('materials','bom','fmd','reports') and r.data.get('material_code')}
        codes.update(code.strip() for code in str(row.data.get('material_codes','')).splitlines() if code.strip())
        result['related']=[serialize(r,limits) for r in supplied if r.module!='materials']
        related_ids=[r.id for r in supplied]
        attachments={}
        if related_ids:
            for evidence in db.scalars(select(Evidence).where(Evidence.record_id.in_(related_ids))).all():
                attachments.setdefault(evidence.record_id,[]).append({'id':evidence.id,'name':evidence.name})
        for item in result['related']:
            item['files']=attachments.get(item['id'],[])
        result['materials']=[serialize(r,limits) for r in visible if r.module=='materials' and (r.data.get('material_code') in codes or (name and r.data.get('supplier')==name))]
    if row.module=='bom':
        code=row.data.get('material_code')
        result['materials']=[serialize(r,limits) for r in all_visible(db,user) if code and r.module=='materials' and r.data.get('material_code')==code]
    return result


class RecordBody(BaseModel):
    module:str
    data:dict
    version:int|None=None


def validate_data(module,data):
    allowed={f['key']:f for f in MODULES[module]['fields']}
    output={}
    for key,value in data.items():
        if key not in allowed: continue
        f=allowed[key]
        if isinstance(value,(dict,list)) or len(str(value))>10000: raise HTTPException(422,'Giá trị không hợp lệ: '+f['label'])
        if f['type']=='select' and value and value not in f['options']: raise HTTPException(422,'Lựa chọn không hợp lệ: '+f['label'])
        if f['type']=='date' and value:
            try: date.fromisoformat(str(value))
            except ValueError: raise HTTPException(422,'Ngày không hợp lệ: '+f['label'])
        if f['type']=='number' and value not in ('',None):
            try:
                value=float(value)
                import math
                if not math.isfinite(value) or value<0: raise ValueError()
            except (ValueError,TypeError): raise HTTPException(422,'Số không hợp lệ: '+f['label'])
        output[key]=value
    for key,f in allowed.items():
        if f['required'] and not str(output.get(key,'')).strip(): raise HTTPException(422,'Thiếu '+f['label'])
    if output.get('expiry_date') and output.get('issue_date') and output['expiry_date']<output['issue_date']: raise HTTPException(422,'Ngày hết hạn phải sau ngày phát hành.')
    return output


def check_workflow(user,db,module,data,record_id=None,before=None):
    before=before or {}
    if module=='xrf-standard': permit(user,db,module,'Approve')
    if data.get('approval') in ('Approved','Rejected') and data.get('approval')!=before.get('approval'):
        permit(user,db,module,'Approve')
    if module=='materials' and data.get('dossier_status')=='Đạt yêu cầu' and before.get('dossier_status')!='Đạt yêu cầu': permit(user,db,module,'Approve')
    if data.get('status')=='Compliant' and before.get('status')!='Compliant': permit(user,db,module,'Approve')
    if data.get('status')=='Closed' and before.get('status')!='Closed':
        permit(user,db,module,'Close')
        if module=='capa':
            count=db.scalar(select(func.count()).select_from(Evidence).where(Evidence.record_id==record_id)) if record_id else 0
            if not all(data.get(k) for k in ('root_cause','corrective','preventive','verification','closure')) or not count:
                raise HTTPException(422,'Đóng CAPA cần root cause, corrective/preventive action, verification, closure và evidence.')


@router.post('/records',status_code=201)
def create(body:RecordBody,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    permit(user,db,body.module,'Create')
    data=validate_data(body.module,body.data); check_workflow(user,db,body.module,data)
    row=Record(module=body.module,data=data); db.add(row); db.flush()
    log(db,user.email,'Tạo hồ sơ',row.id,after=data); db.commit()
    return serialize(row,standards(db))


@router.put('/records/{record_id}')
def edit(record_id:int,body:RecordBody,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    row=record_or_404(db,record_id); permit(user,db,row.module,'Edit')
    if body.module!=row.module: raise HTTPException(422,'Không được đổi module.')
    data=validate_data(row.module,body.data); check_workflow(user,db,row.module,data,row.id,row.data)
    before=row.data
    # Preserve provenance and fields imported from the source which the form does not edit.
    merged={**before,**data}
    result=db.execute(update(Record).where(Record.id==row.id,Record.version==body.version).values(data=merged,version=Record.version+1,updated_at=now()))
    if result.rowcount!=1: db.rollback(); raise HTTPException(409,'Hồ sơ đã được thay đổi. Tải lại trước khi sửa.')
    log(db,user.email,'Cập nhật hồ sơ',row.id,before,merged); db.commit(); db.refresh(row)
    return serialize(row,standards(db))


@router.delete('/records/{record_id}')
def archive(record_id:int,version:int,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    row=record_or_404(db,record_id); permit(user,db,row.module,'Delete')
    changed=db.execute(update(Record).where(Record.id==row.id,Record.version==version).values(archived=True,version=Record.version+1,updated_at=now()))
    if changed.rowcount!=1: db.rollback(); raise HTTPException(409,'Hồ sơ đã thay đổi. Vui lòng tải lại.')
    log(db,user.email,'Lưu trữ hồ sơ',row.id,before=row.data); db.commit(); return {'ok':True}


@router.post('/records/{record_id}/evidence',status_code=201)
async def upload(record_id:int,file:UploadFile=File(...),user:Account=Depends(current_user),db:Session=Depends(get_db)):
    row=record_or_404(db,record_id); permit(user,db,row.module,'Upload')
    name=Path((file.filename or 'file').replace('\\','/')).name[:240]
    ext=Path(name).suffix.lower()
    types={'.pdf':'application/pdf','.png':'image/png','.jpg':'image/jpeg','.jpeg':'image/jpeg','.xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','.txt':'text/plain'}
    if ext not in types: raise HTTPException(422,'Chỉ nhận PDF, PNG, JPG, XLSX, DOCX, TXT.')
    data=await file.read(10*1024*1024+1)
    if not data or len(data)>10*1024*1024: raise HTTPException(413,'File phải từ 1 byte đến 10 MB.')
    signatures={'.pdf':b'%PDF-', '.png':b'\x89PNG\r\n\x1a\n','.jpg':b'\xff\xd8','.jpeg':b'\xff\xd8','.xlsx':b'PK','.docx':b'PK'}
    if ext in signatures and not data.startswith(signatures[ext]): raise HTTPException(422,'Nội dung file không khớp định dạng.')
    UPLOADS.mkdir(parents=True,exist_ok=True); stored=secrets.token_hex(24)+ext; path=UPLOADS/stored
    path.write_bytes(data)
    try:
        e=Evidence(record_id=row.id,name=name,stored_name=stored,mime=types[ext],size=len(data),checksum=hashlib.sha256(data).hexdigest(),uploader=user.name)
        db.add(e); log(db,user.email,'Upload evidence: '+name,row.id,after={'sha256':e.checksum}); db.commit()
    except Exception:
        path.unlink(missing_ok=True); raise
    return {'id':e.id,'name':name}


@router.get('/evidence/{evidence_id}')
def download(evidence_id:int,inline:bool=False,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    e=db.get(Evidence,evidence_id)
    if not e: raise HTTPException(404,'Không tìm thấy file.')
    row=record_or_404(db,e.record_id); permit(user,db,row.module)
    path=UPLOADS/e.stored_name
    if not path.is_file(): raise HTTPException(404,'File evidence không có trên máy chủ.')
    safe_inline=inline and e.mime in ('application/pdf','image/png','image/jpeg')
    return FileResponse(path,media_type=e.mime,filename=e.name,content_disposition_type='inline' if safe_inline else 'attachment',headers={'Content-Security-Policy':"sandbox; default-src 'none'",'X-Content-Type-Options':'nosniff'})


@router.get('/export/{module}')
def export(module:str,q:str='',status:str='',start:str='',end:str='',dossier:str='',user:Account=Depends(current_user),db:Session=Depends(get_db)):
    permit(user,db,module)
    from openpyxl import Workbook
    limits=standards(db)
    rows=filtered([serialize(r,limits) for r in all_visible(db,user)],module,q,status,start,end,dossier=dossier)
    wb=Workbook(); sheet=wb.active; sheet.title='Product Safety'
    columns=MODULES[module]['fields']
    sheet.append([f['label'] for f in columns]+['Trạng thái tính toán','Nguồn file','Sheet','Dòng','Version'])
    for row in rows:
        d=row['data']; source=d.get('_source',{})
        values=[d.get(f['key'],'') for f in columns]+[row['display_status'],source.get('file',''),source.get('sheet',''),source.get('row',''),row['version']]
        sheet.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in values])
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='142B4D')
    for i in range(1,len(columns)+6): sheet.column_dimensions[get_column_letter(i)].width=24
    for row in sheet.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{module}.xlsx"'})


@router.get('/imports')
def imports(user:Account=Depends(current_user),db:Session=Depends(get_db)):
    return {'files':[{'id':r.id,'file':r.filename,'checksum':r.checksum,'summary':r.summary,'at':r.created_at.isoformat()} for r in db.scalars(select(ImportLog)).all()], 'conflicts':workbook_conflicts(all_visible(db,user))}


@router.post('/imports')
def run_import(user:Account=Depends(admin_user),db:Session=Depends(get_db)):
    return import_sources(db,user.email)


@router.get('/imports/{import_id}/download')
def source_download(import_id:int,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    if user.role not in ('Admin','QA Manager','Auditor'): raise HTTPException(403,'Tải workbook nguồn cần Admin, QA Manager hoặc Auditor.')
    item=db.get(ImportLog,import_id)
    if not item or item.filename not in FILES: raise HTTPException(404)
    return FileResponse(ROOT/item.filename,filename=item.filename)


@router.get('/users')
def users(user:Account=Depends(admin_user),db:Session=Depends(get_db)):
    return [profile(u) for u in db.scalars(select(Account).order_by(Account.id)).all()]


class UserUpdate(BaseModel):
    active:bool
    role:str


class ResetPassword(BaseModel):
    password:str=Field(min_length=10,max_length=128)


@router.post('/users/{user_id}/password')
def reset_password(user_id:int,body:ResetPassword,user:Account=Depends(admin_user),db:Session=Depends(get_db)):
    from login import password_hash
    target=db.get(Account,user_id)
    if not target: raise HTTPException(404)
    target.password_hash=password_hash(body.password)
    target.failed=0; target.locked_until=None
    db.execute(delete(LoginSession).where(LoginSession.user_id==target.id))
    log(db,user.email,'Admin đặt lại mật khẩu',after={'account_id':target.id})
    db.commit(); return {'ok':True}


@router.put('/users/{user_id}')
def update_user(user_id:int,body:UserUpdate,user:Account=Depends(admin_user),db:Session=Depends(get_db)):
    target=db.get(Account,user_id)
    if not target: raise HTTPException(404)
    if body.role not in ROLES: raise HTTPException(422,'Role không hợp lệ.')
    if target.id==user.id and (not body.active or body.role!='Admin'): raise HTTPException(422,'Không được tự khóa hoặc hạ quyền Admin đang đăng nhập.')
    before=profile(target); target.active=body.active; target.role=body.role
    db.execute(delete(LoginSession).where(LoginSession.user_id==target.id))
    log(db,user.email,'Cập nhật quyền tài khoản',before=before,after=profile(target)); db.commit(); return profile(target)


@router.get('/settings/{key}')
def setting(key:str,user:Account=Depends(current_user),db:Session=Depends(get_db)):
    if key not in ('permissions','notifications','retention','system'): raise HTTPException(404)
    row=db.get(Setting,key)
    if row: return row.value
    if key=='notifications': return {'days':30,'reports':True,'training':True,'capa':True,'documents':True,'compliance':True}
    if key=='retention': return {'documents':5,'samples':2,'reports':5,'audit':5,'training':3}
    if key=='system': return {'site_name':'NFV2 Product Safety','factory':'NFV2','timezone':'Asia/Bangkok'}
    return {role:permissions(Account(role=role),db) for role in ROLES if role!='Admin'}


@router.put('/settings/{key}')
def save_setting(key:str,body:dict,user:Account=Depends(admin_user),db:Session=Depends(get_db)):
    if key not in ('permissions','notifications','retention','system'): raise HTTPException(404)
    if key=='permissions':
        if any(role not in ROLES or role=='Admin' for role in body): raise HTTPException(422,'Không được thay đổi quyền Admin.')
        for role,modules in body.items():
            if not isinstance(modules,dict): raise HTTPException(422)
            for mod,actions in modules.items():
                if mod not in MODULES or not isinstance(actions,dict) or any(a not in ACTIONS or not isinstance(v,bool) for a,v in actions.items()): raise HTTPException(422)
    elif key=='notifications':
        if body.get('days') not in (30,60,90): raise HTTPException(422,'Chọn 30, 60 hoặc 90 ngày.')
        if any(not isinstance(v,bool) for k,v in body.items() if k!='days'): raise HTTPException(422)
    elif key=='retention':
        if any(not isinstance(v,int) or not 1<=v<=100 for v in body.values()): raise HTTPException(422,'Thời gian lưu: 1–100 năm.')
    elif any(not isinstance(v,str) or len(v)>200 for v in body.values()): raise HTTPException(422)
    row=db.get(Setting,key); before=row.value if row else {}
    if row: row.value=body
    else: db.add(Setting(key=key,value=body))
    log(db,user.email,'Cài đặt: '+key,before=before,after=body); db.commit(); return body


@router.post('/bom/import')
def import_bom_api(user:Account=Depends(current_user), db:Session=Depends(get_db)):
    permit(user, db, 'bom')
    from backend.excel_import import import_project_bom
    res = import_project_bom(db)
    if res.get('status') == 'error':
        raise HTTPException(400, res.get('message', 'Lỗi nhập BOM'))
    log(db, user.email, 'Cập nhật BOM dự án từ file Excel', after=res)
    return res


@router.post('/bom/save')
def save_bom_api(user:Account=Depends(current_user), db:Session=Depends(get_db)):
    permit(user, db, 'bom', 'Edit')
    count = db.query(Record).filter(Record.module=='bom', Record.archived==False).count()
    log(db, user.email, 'Lưu và đồng bộ dữ liệu BOM', after={'total_bom': count})
    db.commit()
    return {'status':'ok', 'message':f'Đã lưu thành công {count} bản ghi BOM.', 'total':count}


@router.delete('/bom')
def delete_bom_api(project:str='', user:Account=Depends(current_user), db:Session=Depends(get_db)):
    permit(user, db, 'bom', 'Delete')
    query = db.query(Record).filter(Record.module=='bom', Record.archived==False)
    if project:
        rows = [r for r in query.all() if r.data.get('project')==project or r.data.get('parent_code')==project]
    else:
        rows = query.all()
    count = len(rows)
    deleted_codes = {r.data.get('material_code') for r in rows if r.data.get('material_code')}
    deleted_projects = {r.data.get('project') for r in rows if r.data.get('project')}
    if project:
        deleted_projects.add(project)

    for r in rows:
        r.archived = True
        r.updated_at = now()

    # Determine remaining BOM codes and suppliers
    remaining_bom_records = db.query(Record).filter(Record.module=='bom', Record.archived==False).all()
    remaining_bom_codes = {r.data.get('material_code') for r in remaining_bom_records if r.data.get('material_code')}
    remaining_suppliers = {r.data.get('supplier') for r in remaining_bom_records if r.data.get('supplier')}

    # Synchronize materials: archive materials belonging to deleted BOM
    mat_query = db.query(Record).filter(Record.module=='materials', Record.archived==False)
    mats_archived = 0
    for m in mat_query.all():
        code = m.data.get('material_code')
        mat_proj = m.data.get('project', '')
        if project:
            if mat_proj == project or (code in deleted_codes and code not in remaining_bom_codes):
                m.archived = True
                m.updated_at = now()
                mats_archived += 1
        else:
            if code in deleted_codes or mat_proj in deleted_projects:
                m.archived = True
                m.updated_at = now()
                mats_archived += 1

    # Synchronize suppliers: archive suppliers not present in remaining BOM
    sup_query = db.query(Record).filter(Record.module=='suppliers', Record.archived==False)
    sups_archived = 0
    for s in sup_query.all():
        sup_name = s.data.get('supplier')
        if sup_name not in remaining_suppliers:
            s.archived = True
            s.updated_at = now()
            sups_archived += 1

    log(db, user.email, f'Xóa dữ liệu BOM ({project if project else "Toàn bộ"}) & đồng bộ', after={'bom_count': count, 'materials_count': mats_archived, 'suppliers_count': sups_archived})
    db.commit()
    return {'status':'ok', 'count':count, 'materials_count': mats_archived, 'suppliers_count': sups_archived}



POLYMERS_DEFAULTS = [
    {'element':'Cd (Cadmium)','material_type':'Polymers','control_limit':35,'spec_limit':100,'rule':'≤ Control Limit'},
    {'element':'Pb (Lead)','material_type':'Polymers','control_limit':35,'spec_limit':1000,'rule':'≤ Control Limit'},
    {'element':'Hg (Mercury)','material_type':'Polymers','control_limit':70,'spec_limit':1000,'rule':'≤ Control Limit'},
    {'element':'Cr (Chromium)','material_type':'Polymers','control_limit':350,'spec_limit':1000,'rule':'≤ Control Limit'},
    {'element':'Br (Bromine)','material_type':'Polymers','control_limit':300,'spec_limit':1000,'rule':'≤ Control Limit'},
    {'element':'Cl (Chlorine)','material_type':'Polymers','control_limit':630,'spec_limit':1500,'rule':'≤ Control Limit'},
]


@router.get('/xrf-limits')
def xrf_limits(user:Account=Depends(current_user), db:Session=Depends(get_db)):
    permit(user, db, 'xrf-standard')
    rows = db.scalars(select(Record).where(Record.module=='xrf-standard', Record.archived==False)).all()
    limits = standards(db)
    return [serialize(r, limits) for r in rows]


@router.post('/xrf-limits/seed')
def seed_xrf_limits(user:Account=Depends(current_user), db:Session=Depends(get_db)):
    permit(user, db, 'xrf-standard', 'Approve')
    existing = db.scalars(select(Record).where(Record.module=='xrf-standard', Record.archived==False)).all()
    existing_keys = {(str(r.data.get('element','')), str(r.data.get('material_type',''))) for r in existing}
    created = 0
    for item in POLYMERS_DEFAULTS:
        key = (item['element'], item['material_type'])
        if key not in existing_keys:
            db.add(Record(module='xrf-standard', data=item, version=1, updated_at=now()))
            created += 1
    if created:
        log(db, user.email, f'Seed {created} XRF control limits (Polymers)', after={'created': created})
        db.commit()
    return {'created': created, 'total': len(existing) + created}

